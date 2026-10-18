#!/usr/bin/env python3
"""
D&D Tribe Excel -> TypeScript Sync
==================================

Purpose
-------
Use "Put into website as tribes.xlsx" as the canonical item inventory and
sync the items into your EXISTING tribe*.ts files.

This v2 version prefers an explicit shop/tribe -> REAL filename map based on
the filenames you supplied. It never invents a new tribe filename.

For each spreadsheet item the program writes/updates:

    name
    rarity
    itemType
    attunment
    price
    description

It preserves tribe-file-only fields such as:

    priceLabel
    priceText
    or any other custom item property already present

It also preserves the rest of each tribe file:
owners, insults, custom interfaces, price variability, etc.

IMPORTANT
---------
The spreadsheet is treated as authoritative for:
    rarity, item type, attunment, price, and substantive description changes.

Existing TypeScript item names are preserved when an Excel row is matched to
an existing item. This matters for examples such as:
    Excel: Amber
    TS:    Amber (Kalashtar, Psychic human)

New Excel-only ITEMS are added to the matched tribe's items array.

New tribe FILES are never invented. If a spreadsheet tribe cannot be matched
to one of your existing .ts files, it is reported and skipped.

Existing TS-only items are PRESERVED at the bottom of the items array by
default and reported for review. Use --drop-ts-only only if you want the
spreadsheet to be the exact source of truth.

The program is conservative about fuzzy matching and writes a detailed report.

Dependencies
------------
    python -m pip install openpyxl

Typical use
-----------
Put this script, the workbook, and the tribe*.ts files in the same project
folder, then run:

    python dnd_tribe_excel_to_ts_sync.py

Or specify locations:

    python dnd_tribe_excel_to_ts_sync.py \
        "E:\\Put into website as tribes.xlsx" \
        --tribes-dir "E:\\my-project\\src\\tribes"

By default it creates:
    tribes_updated_from_excel/

and does NOT overwrite the originals.

To overwrite the original tribe files with backups:
    python dnd_tribe_excel_to_ts_sync.py --in-place

Dry run:
    python dnd_tribe_excel_to_ts_sync.py --dry-run

If your central types.ts Item interface does not yet contain the new fields:
    python dnd_tribe_excel_to_ts_sync.py --patch-types

That adds optional fields:
    rarity?: string;
    itemType?: string;
    attunment?: boolean;

to interface Item / type Item = { ... } when types.ts can be found.
"""

import argparse
import json
import os
import re
import shutil
import sys
import textwrap
import unicodedata
from dataclasses import dataclass
from difflib import SequenceMatcher
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from openpyxl import load_workbook


# ---------------------------------------------------------------------------
# CONFIG
# ---------------------------------------------------------------------------

DEFAULT_WORKBOOK = "Put into website as tribes.xlsx"
DEFAULT_SHEET = "Combination"

HEADERS = {
    "name": "Item Name",
    "rarity": "Rarity",
    "item_type": "Item Type",
    "attunment": "Attunment",
    "description": "Description",
    "price": "Price",
    "shop": "What shop it should go to",
}

KNOWN_SYNC_FIELDS = {
    "name",
    "rarity",
    "itemType",
    "attunment",
    "attunement",  # remove old correctly-spelled variant if present
    "price",
    "description",
}

OUTPUT_FOLDER_NAME = "tribes_updated_from_excel"
REPORT_NAME = "tribe_sync_report.txt"

FUZZY_THRESHOLD = 0.84
PREFIX_THRESHOLD = 0.90
SMART_DESCRIPTION_SIMILARITY = 0.78


# ---------------------------------------------------------------------------
# REAL TRIBE FILENAMES
# ---------------------------------------------------------------------------
#
# These are the actual filenames you showed me. The LEFT side is the shop /
# tribe name coming from Excel. The RIGHT side is the existing TypeScript file.
#
# IMPORTANT:
# - The program checks this map FIRST.
# - If the mapped file exists, it uses that exact file.
# - If a tribe is not in this map, the program may still find it by reading
#   the top-level `name:` inside an existing tribe*.ts file (useful for
#   tribe1.ts, tribe2.ts, etc.).
# - If neither method finds an existing file, the tribe is skipped and
#   reported. The program DOES NOT make up a new filename.
#
TRIBE_FILE_MAP = {
    "Applegarth Guild": "tribeApplegarthGuild.ts",
    "Archives Guild": "tribeArchivesGuild.ts",
    "Auntie Patty's Pies": "tribeAuntiePattysPies.ts",
    "Blossom Hotel": "tribeBlossomHotel.ts",
    "Book Bombs": "tribeBookBombs.ts",
    "Bullets, Buffs, & Beyond": "tribeBulletsBuffsBeyond.ts",
    "Changing Church": "tribeChangingChurch.ts",
    "Comedy Gold": "tribeComedyGold.ts",
    "Dungeon Crawler Guild": "tribeDungeonCrawlerGuild.ts",
    "Evan's Enchanting Emporium": "tribeEvansEnchantingEmporium.ts",
    "Fairies of Flora": "tribeFairiesOfFlora.ts",
    "Find a Friend": "tribeFindAFriend.ts",
    "Fizzy Tales": "tribeFizzyTales.ts",
    "Golem Workshop": "tribeGolemWorkshop.ts",
    "Iconic Dragonic": "tribeIconicDragonic.ts",
    "Jell Bell": "tribeJellBell.ts",
    "Jewelry Guild": "tribeJewelryGuild.ts",
    "Labyrinthine Library": "tribeLabyrinthineLibrary.ts",
    "Make a Monster": "tribeMonsterMaker.ts",
    "Michael's Mount": "tribeMichaelsMount.ts",
    "N.M.E.": "tribeNME.ts",
    "Navigation Guild": "tribeNavigationGuild.ts",
    "Necromancy Insurance Company": "tribeNecromancyInsuranceCompany.ts",
    "O-Papies Oracle Readings": "tribeOPapiesOracleReadings.ts",
    "Paws, Claws, & Maws": "tribePawsClawsMaws.ts",
    "Pearl's Potions": "tribePearlsPotions.ts",
    "Provision's Paradise": "tribeProvisionsParadise.ts",
    "Robin's Ropes": "tribeRobinsRopes.ts",
    "Runestone Relay": "tribeRunestoneRelay.ts",
    "Silent Oath": "tribeSilentOath.ts",
    "Sleuth University": "tribeSleuthUniversity.ts",
    "Supreme Smithy": "tribeSupremeSmithy.ts",
    "The Piggy Bank, no hammers inside.": "tribePiggyBank.ts",
    "Valhalla Mart": "tribeValhallaMart.ts",
    "Will's Weapons": "tribeWillsWeapons.ts",
    "Ye Old Donkey": "tribeYeOldDonkey.ts",
    "Ye Old Home Depot": "tribeYeOldHomeDepot.ts",
}

# These files were also visible in your folder. They are NOT automatically
# assigned to a shop name because the screenshot alone does not tell us what
# their top-level tribe names are. The fallback discovery code reads their
# actual `name:` field, so they can still be used safely when appropriate.
KNOWN_EXISTING_TRIBE_FILENAMES = {
    "tribeAuctionHouse.ts",
    "tribeAuctionHouse2.ts",
    "tribeAuctionHouse3.ts",
    "tribeAuctionHouse4.ts",
    "tribeAuctionHouse5.ts",
    "tribeBlackMarket.ts",
    "tribeBlackMarket2.ts",
    "tribeBlackMarket3.ts",
    "tribeBlackMarket4.ts",
    "tribeBlackMarket5.ts",
    "tribeInsults.ts",
    "tribe1.ts",
    "tribe2.ts",
    "tribe3.ts",
    "tribe4.ts",
    "tribe5.ts",
    "tribe6.ts",
    "tribe7.ts",
}


# ---------------------------------------------------------------------------
# DATA MODELS
# ---------------------------------------------------------------------------

@dataclass
class ExcelItem:
    row: int
    name: str
    rarity: str
    item_type: str
    attunment: bool
    description: str
    price: Optional[float]
    shop_label: str
    tribe_name: str


@dataclass
class ExistingItem:
    index: int
    object_text: str
    properties: List[Tuple[str, str]]
    property_map: Dict[str, str]
    name: str


@dataclass
class MatchResult:
    existing: Optional[ExistingItem]
    match_type: str
    score: float


# ---------------------------------------------------------------------------
# TEXT NORMALIZATION
# ---------------------------------------------------------------------------

def clean_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def normalize_text(value: str) -> str:
    text = unicodedata.normalize("NFKC", clean_text(value))
    text = (
        text.replace("’", "'")
        .replace("‘", "'")
        .replace("“", '"')
        .replace("”", '"')
        .replace("–", "-")
        .replace("—", "-")
        .lower()
    )
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def normalize_name(value: str) -> str:
    text = normalize_text(value)
    text = re.sub(r"\s*,\s*", ", ", text)
    return text.strip(" .")


def strip_trailing_parentheticals(value: str) -> str:
    """
    Remove one or more trailing parenthetical notes from a display name.

    Examples:
        Amber (Kalashtar, Psychic human) -> Amber
        Hire X (Local) for One Job       -> unchanged by this function
    """
    text = clean_text(value).strip()

    while True:
        updated = re.sub(r"\s*\([^()]*\)\s*$", "", text).strip()
        if updated == text:
            break
        text = updated

    return text


def compact_name_for_match(value: str) -> str:
    text = strip_trailing_parentheticals(value)

    # Also remove punctuation that commonly differs between Excel/TS labels.
    text = normalize_text(text)
    text = re.sub(r"[\[\]{}(),:;.!?]", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def descriptions_similar(left: str, right: str) -> float:
    a = normalize_text(left)
    b = normalize_text(right)

    # Formatting punctuation / paragraph breaks should not dominate.
    a = re.sub(r"[^a-z0-9'+\-]+", " ", a)
    b = re.sub(r"[^a-z0-9'+\-]+", " ", b)
    a = re.sub(r"\s+", " ", a).strip()
    b = re.sub(r"\s+", " ", b).strip()

    if not a or not b:
        return 0.0

    return SequenceMatcher(None, a, b).ratio()


# ---------------------------------------------------------------------------
# EXCEL
# ---------------------------------------------------------------------------

def tribe_name_from_shop_label(label: str) -> str:
    """
    Spreadsheet labels look like:
        Applegarth Guild (Companions, Hirelings)
        Book Bombs (Specialty, Book store)

    The tribe TS object uses only:
        Applegarth Guild
        Book Bombs
    """
    text = clean_text(label)

    # Only split on " (" so punctuation in names is preserved.
    if " (" in text:
        return text.split(" (", 1)[0].strip()

    return text.strip()


def parse_bool(value) -> bool:
    if isinstance(value, bool):
        return value

    text = normalize_text(value)

    if text in {"true", "yes", "y", "1"}:
        return True

    return False


def parse_price(value) -> Optional[float]:
    if value is None or clean_text(value) == "":
        return None

    if isinstance(value, (int, float)):
        return float(value)

    text = clean_text(value).replace(",", "").replace("$", "")

    try:
        return float(text)
    except ValueError:
        return None


def read_excel_items(
    workbook_path: Path,
    sheet_name: str,
) -> List[ExcelItem]:
    wb = load_workbook(
        workbook_path,
        read_only=True,
        data_only=True,
    )

    if sheet_name not in wb.sheetnames:
        raise ValueError(
            f'Sheet "{sheet_name}" was not found. '
            f"Available: {', '.join(wb.sheetnames)}"
        )

    ws = wb[sheet_name]

    header_row = next(
        ws.iter_rows(
            min_row=1,
            max_row=1,
            values_only=True,
        )
    )

    positions: Dict[str, int] = {}

    for index, value in enumerate(header_row):
        header = clean_text(value)

        for key, wanted in HEADERS.items():
            if header == wanted and key not in positions:
                positions[key] = index

    missing = [
        wanted
        for key, wanted in HEADERS.items()
        if key not in positions
    ]

    if missing:
        raise ValueError(
            "Missing spreadsheet header(s): "
            + ", ".join(missing)
        )

    items: List[ExcelItem] = []

    for excel_row_number, row in enumerate(
        ws.iter_rows(
            min_row=2,
            values_only=True,
        ),
        start=2,
    ):
        name = clean_text(
            row[positions["name"]]
            if positions["name"] < len(row)
            else None
        )

        shop_label = clean_text(
            row[positions["shop"]]
            if positions["shop"] < len(row)
            else None
        )

        if not name or not shop_label:
            continue

        rarity = clean_text(
            row[positions["rarity"]]
            if positions["rarity"] < len(row)
            else None
        )

        item_type = clean_text(
            row[positions["item_type"]]
            if positions["item_type"] < len(row)
            else None
        )

        description = clean_text(
            row[positions["description"]]
            if positions["description"] < len(row)
            else None
        )

        attunment = parse_bool(
            row[positions["attunment"]]
            if positions["attunment"] < len(row)
            else None
        )

        price = parse_price(
            row[positions["price"]]
            if positions["price"] < len(row)
            else None
        )

        items.append(
            ExcelItem(
                row=excel_row_number,
                name=name,
                rarity=rarity,
                item_type=item_type,
                attunment=attunment,
                description=description,
                price=price,
                shop_label=shop_label,
                tribe_name=tribe_name_from_shop_label(
                    shop_label
                ),
            )
        )

    return items


# ---------------------------------------------------------------------------
# TYPESCRIPT SCANNER
# ---------------------------------------------------------------------------

def skip_string_or_comment(
    text: str,
    index: int,
) -> Optional[int]:
    """
    If text[index] starts a JS/TS string or comment, return the index after it.
    Otherwise return None.
    """
    if index >= len(text):
        return None

    ch = text[index]

    # Quoted strings and template literals.
    if ch in {'"', "'", "`"}:
        quote = ch
        i = index + 1

        while i < len(text):
            if text[i] == "\\":
                i += 2
                continue

            if text[i] == quote:
                return i + 1

            i += 1

        return len(text)

    # Line comment.
    if (
        ch == "/"
        and index + 1 < len(text)
        and text[index + 1] == "/"
    ):
        newline = text.find("\n", index + 2)
        return len(text) if newline == -1 else newline + 1

    # Block comment.
    if (
        ch == "/"
        and index + 1 < len(text)
        and text[index + 1] == "*"
    ):
        end = text.find("*/", index + 2)
        return len(text) if end == -1 else end + 2

    return None


def find_matching_delimiter(
    text: str,
    open_index: int,
    open_char: str,
    close_char: str,
) -> int:
    depth = 0
    i = open_index

    while i < len(text):
        skipped = skip_string_or_comment(
            text,
            i,
        )

        if skipped is not None:
            i = skipped
            continue

        ch = text[i]

        if ch == open_char:
            depth += 1

        elif ch == close_char:
            depth -= 1

            if depth == 0:
                return i

        i += 1

    raise ValueError(
        f"Could not find matching {close_char} "
        f"for {open_char} at character {open_index}."
    )


def find_items_array_region(
    source: str,
) -> Tuple[int, int, int, int]:
    match = re.search(
        r"\bitems\s*:\s*\[",
        source,
    )

    if not match:
        raise ValueError(
            'Could not find an "items: [" array.'
        )

    open_bracket = source.find(
        "[",
        match.start(),
        match.end(),
    )

    close_bracket = find_matching_delimiter(
        source,
        open_bracket,
        "[",
        "]",
    )

    return (
        match.start(),
        open_bracket,
        close_bracket,
        close_bracket + 1,
    )


def split_top_level_item_objects(
    array_content: str,
) -> List[str]:
    objects: List[str] = []
    i = 0

    while i < len(array_content):
        skipped = skip_string_or_comment(
            array_content,
            i,
        )

        if skipped is not None:
            i = skipped
            continue

        if array_content[i] == "{":
            end = find_matching_delimiter(
                array_content,
                i,
                "{",
                "}",
            )

            objects.append(
                array_content[i:end + 1]
            )

            i = end + 1
            continue

        i += 1

    return objects


def read_property_key(
    text: str,
    index: int,
) -> Tuple[Optional[str], int]:
    i = index

    while (
        i < len(text)
        and text[i].isspace()
    ):
        i += 1

    if i >= len(text):
        return None, i

    # Identifier key.
    match = re.match(
        r"[A-Za-z_$][A-Za-z0-9_$]*",
        text[i:],
    )

    if match:
        key = match.group(0)
        return key, i + len(key)

    # Quoted key.
    if text[i] in {'"', "'"}:
        quote = text[i]
        end = i + 1

        while end < len(text):
            if text[end] == "\\":
                end += 2
                continue

            if text[end] == quote:
                raw = text[i:end + 1]
                return decode_ts_string(raw), end + 1

            end += 1

    return None, i


def find_top_level_value_end(
    text: str,
    start: int,
    object_end: int,
) -> int:
    """
    Find the comma ending one top-level object property.
    """
    i = start

    paren = 0
    bracket = 0
    brace = 0

    while i < object_end:
        skipped = skip_string_or_comment(
            text,
            i,
        )

        if skipped is not None:
            i = skipped
            continue

        ch = text[i]

        if ch == "(":
            paren += 1
        elif ch == ")":
            paren = max(0, paren - 1)
        elif ch == "[":
            bracket += 1
        elif ch == "]":
            bracket = max(0, bracket - 1)
        elif ch == "{":
            brace += 1
        elif ch == "}":
            if brace > 0:
                brace -= 1
        elif (
            ch == ","
            and paren == 0
            and bracket == 0
            and brace == 0
        ):
            return i

        i += 1

    return object_end


def parse_object_properties(
    object_text: str,
) -> List[Tuple[str, str]]:
    """
    Parse top-level properties from one object literal.

    Returns:
        [(key, raw_value_text), ...]

    Unknown/custom values are preserved as raw TypeScript.
    """
    if not object_text.strip().startswith("{"):
        return []

    object_end = len(object_text) - 1
    i = 1
    properties: List[Tuple[str, str]] = []

    while i < object_end:
        # Skip whitespace and commas.
        while (
            i < object_end
            and (
                object_text[i].isspace()
                or object_text[i] == ","
            )
        ):
            i += 1

        if i >= object_end:
            break

        # Skip comments.
        skipped = skip_string_or_comment(
            object_text,
            i,
        )

        if skipped is not None and object_text[i] == "/":
            i = skipped
            continue

        key, after_key = read_property_key(
            object_text,
            i,
        )

        if not key:
            # Unknown syntax. Advance one char rather than looping.
            i += 1
            continue

        i = after_key

        while (
            i < object_end
            and object_text[i].isspace()
        ):
            i += 1

        if (
            i >= object_end
            or object_text[i] != ":"
        ):
            # Method/shorthand/spread syntax is not expected in item objects.
            # Leave it unparsed rather than guessing.
            i += 1
            continue

        i += 1

        while (
            i < object_end
            and object_text[i].isspace()
        ):
            i += 1

        value_start = i
        value_end = find_top_level_value_end(
            object_text,
            value_start,
            object_end,
        )

        raw_value = object_text[
            value_start:value_end
        ].strip()

        properties.append(
            (
                key,
                raw_value,
            )
        )

        i = (
            value_end + 1
            if value_end < object_end
            else object_end
        )

    return properties


def decode_ts_string(raw: str) -> str:
    raw = clean_text(raw)

    if len(raw) < 2:
        return raw

    if raw[0] == '"' and raw[-1] == '"':
        try:
            return json.loads(raw)
        except Exception:
            return raw[1:-1]

    if raw[0] == "'" and raw[-1] == "'":
        body = raw[1:-1]
        body = body.replace(r"\'", "'")
        body = body.replace(r"\\", "\\")
        body = body.replace(r"\n", "\n")
        body = body.replace(r"\r", "\r")
        body = body.replace(r"\t", "\t")
        return body

    if raw[0] == "`" and raw[-1] == "`":
        body = raw[1:-1]
        body = body.replace(r"\`", "`")
        body = body.replace(r"\${", "${")
        return body

    return raw


def parse_existing_items(
    source: str,
) -> List[ExistingItem]:
    _, open_bracket, close_bracket, _ = (
        find_items_array_region(source)
    )

    array_content = source[
        open_bracket + 1:close_bracket
    ]

    object_texts = split_top_level_item_objects(
        array_content
    )

    result: List[ExistingItem] = []

    for index, object_text in enumerate(
        object_texts
    ):
        props = parse_object_properties(
            object_text
        )

        prop_map = {
            key: raw
            for key, raw in props
        }

        raw_name = prop_map.get(
            "name",
            "",
        )

        name = decode_ts_string(
            raw_name
        )

        result.append(
            ExistingItem(
                index=index,
                object_text=object_text,
                properties=props,
                property_map=prop_map,
                name=name,
            )
        )

    return result


def extract_tribe_name(
    source: str,
) -> str:
    """
    Find the tribe's top-level name before items:.
    """
    items_match = re.search(
        r"\bitems\s*:\s*\[",
        source,
    )

    prefix = (
        source[:items_match.start()]
        if items_match
        else source
    )

    matches = list(
        re.finditer(
            r"\bname\s*:\s*(\"(?:\\.|[^\"\\])*\"|'(?:\\.|[^'\\])*')",
            prefix,
        )
    )

    if not matches:
        return ""

    return decode_ts_string(
        matches[-1].group(1)
    )


# ---------------------------------------------------------------------------
# MATCHING
# ---------------------------------------------------------------------------

def score_item_name_match(
    excel_name: str,
    ts_name: str,
) -> Tuple[float, str]:
    a = normalize_name(
        excel_name
    )
    b = normalize_name(
        ts_name
    )

    if not a or not b:
        return 0.0, "none"

    if a == b:
        return 1.0, "exact"

    compact_a = compact_name_for_match(
        excel_name
    )
    compact_b = compact_name_for_match(
        ts_name
    )

    if (
        compact_a
        and compact_a == compact_b
    ):
        return 0.985, "parenthetical-normalized"

    # A shorter spreadsheet label often becomes a richer website display name.
    if (
        len(compact_a) >= 5
        and (
            compact_b.startswith(
                compact_a + " "
            )
            or compact_a.startswith(
                compact_b + " "
            )
        )
    ):
        return 0.94, "prefix"

    # Handle cases like:
    #   Hire an Apology Courier
    #   Hire an Apology Courier (Local) for One Job
    if (
        len(a) >= 8
        and (
            b.startswith(a)
            or a.startswith(b)
        )
    ):
        return 0.93, "display-prefix"

    ratio = SequenceMatcher(
        None,
        compact_a,
        compact_b,
    ).ratio()

    if ratio >= FUZZY_THRESHOLD:
        return ratio, "fuzzy"

    return ratio, "none"


def find_best_existing_match(
    excel_item: ExcelItem,
    existing_items: List[ExistingItem],
    already_used: set,
) -> MatchResult:
    best: Optional[ExistingItem] = None
    best_score = 0.0
    best_type = "none"

    for item in existing_items:
        if item.index in already_used:
            continue

        score, match_type = (
            score_item_name_match(
                excel_item.name,
                item.name,
            )
        )

        if score > best_score:
            best = item
            best_score = score
            best_type = match_type

    if best is None:
        return MatchResult(
            existing=None,
            match_type="new",
            score=0.0,
        )

    minimum = (
        FUZZY_THRESHOLD
        if best_type == "fuzzy"
        else PREFIX_THRESHOLD
        if best_type in {
            "prefix",
            "display-prefix",
        }
        else 0.97
    )

    if best_score < minimum:
        return MatchResult(
            existing=None,
            match_type="new",
            score=best_score,
        )

    return MatchResult(
        existing=best,
        match_type=best_type,
        score=best_score,
    )


# ---------------------------------------------------------------------------
# DESCRIPTION SMART UPDATE
# ---------------------------------------------------------------------------

def initial_parenthetical_body(
    description: str,
) -> Optional[Tuple[str, str]]:
    """
    Split:
        (Core description): Display label

    while respecting inner parentheses.
    """
    text = clean_text(description)

    if not text.startswith("("):
        return None

    depth = 0

    for index, ch in enumerate(text):
        if ch == "(":
            depth += 1

        elif ch == ")":
            depth -= 1

            if depth == 0:
                remainder = text[
                    index + 1:
                ].strip()

                if remainder.startswith(":"):
                    return (
                        text[1:index].strip(),
                        remainder[1:].strip(),
                    )

                return None

    return None


def choose_description_value(
    excel_description: str,
    existing: Optional[ExistingItem],
) -> Tuple[str, bool, str]:
    """
    Returns:
        raw TS value
        preserved_existing_format
        decision_reason

    Smart mode preserves an existing nicely formatted description when Excel
    is materially the same information. If Excel changed the substance, Excel
    wins.
    """
    excel_description = clean_text(
        excel_description
    )

    if existing is None:
        return (
            ts_string(
                excel_description
            ),
            False,
            "new item: Excel description",
        )

    raw_existing = existing.property_map.get(
        "description",
        "",
    )

    existing_text = decode_ts_string(
        raw_existing
    )

    # Comedy Gold-style rows may combine:
    #   (description): priceLabel
    # in Excel, while TS correctly separates description and priceLabel.
    if (
        "priceLabel"
        in existing.property_map
    ):
        split = initial_parenthetical_body(
            excel_description
        )

        if split is not None:
            body, _label = split
            excel_description = body

    similarity = descriptions_similar(
        excel_description,
        existing_text,
    )

    if (
        raw_existing
        and similarity
        >= SMART_DESCRIPTION_SIMILARITY
    ):
        return (
            raw_existing,
            True,
            (
                "kept existing TS formatting "
                f"(description similarity {similarity:.3f})"
            ),
        )

    return (
        ts_string(
            excel_description
        ),
        False,
        (
            "updated from Excel "
            f"(description similarity {similarity:.3f})"
        ),
    )


# ---------------------------------------------------------------------------
# TYPESCRIPT OUTPUT
# ---------------------------------------------------------------------------

def ts_string(value: str) -> str:
    text = clean_text(value)

    # Preserve real line breaks as readable template literals.
    if "\n" in text or "\r" in text:
        escaped = (
            text.replace("\\", "\\\\")
            .replace("`", r"\`")
            .replace("${", r"\${")
        )
        return "`" + escaped + "`"

    return json.dumps(
        text,
        ensure_ascii=False,
    )


def ts_number(value: Optional[float]) -> str:
    if value is None:
        return "0"

    if float(value).is_integer():
        return str(
            int(value)
        )

    return str(
        round(
            float(value),
            6,
        )
    ).rstrip("0").rstrip(".")


def build_item_object(
    excel_item: ExcelItem,
    existing: Optional[ExistingItem],
) -> Tuple[str, str]:
    """
    Build a clean object literal using Excel's canonical synced fields and
    preserving unknown/custom TS properties.
    """
    # Preserve website display name if this is an existing matched item.
    if (
        existing is not None
        and existing.property_map.get(
            "name"
        )
    ):
        name_raw = existing.property_map[
            "name"
        ]
    else:
        name_raw = ts_string(
            excel_item.name
        )

    description_raw, _, desc_reason = (
        choose_description_value(
            excel_item.description,
            existing,
        )
    )

    properties: List[
        Tuple[str, str]
    ] = [
        (
            "name",
            name_raw,
        ),
        (
            "rarity",
            ts_string(
                excel_item.rarity
            ),
        ),
        (
            "itemType",
            ts_string(
                excel_item.item_type
            ),
        ),
        (
            "attunment",
            (
                "true"
                if excel_item.attunment
                else "false"
            ),
        ),
        (
            "price",
            ts_number(
                excel_item.price
            ),
        ),
        (
            "description",
            description_raw,
        ),
    ]

    # Preserve custom fields from the old object.
    if existing is not None:
        for key, raw_value in existing.properties:
            if key in KNOWN_SYNC_FIELDS:
                continue

            properties.append(
                (
                    key,
                    raw_value,
                )
            )

    lines = [
        "    {",
    ]

    for key, raw_value in properties:
        # Multiline template literals need indentation on property start only.
        lines.append(
            f"      {key}: {raw_value},"
        )

    lines.append(
        "    }"
    )

    return (
        "\n".join(
            lines
        ),
        desc_reason,
    )


def rewrite_items_array(
    source: str,
    rebuilt_objects: List[str],
) -> str:
    (
        _items_start,
        open_bracket,
        close_bracket,
        _end,
    ) = find_items_array_region(
        source
    )

    if rebuilt_objects:
        replacement = (
            "[\n"
            + ",\n".join(
                rebuilt_objects
            )
            + "\n  ]"
        )
    else:
        replacement = "[]"

    return (
        source[:open_bracket]
        + replacement
        + source[close_bracket + 1:]
    )


# ---------------------------------------------------------------------------
# TYPES.TS PATCH
# ---------------------------------------------------------------------------

def patch_item_interface(
    source: str,
) -> Tuple[str, bool, str]:
    needed = [
        (
            "rarity",
            "  rarity?: string;",
        ),
        (
            "itemType",
            "  itemType?: string;",
        ),
        (
            "attunment",
            "  attunment?: boolean;",
        ),
    ]

    missing = [
        line
        for key, line in needed
        if not re.search(
            rf"\b{re.escape(key)}\s*\??\s*:",
            source,
        )
    ]

    if not missing:
        return (
            source,
            False,
            "types.ts already contains rarity/itemType/attunment",
        )

    patterns = [
        r"(?:export\s+)?interface\s+Item\s*\{",
        r"(?:export\s+)?type\s+Item\s*=\s*\{",
    ]

    match = None

    for pattern in patterns:
        match = re.search(
            pattern,
            source,
        )

        if match:
            break

    if not match:
        return (
            source,
            False,
            "could not locate interface/type Item in types.ts",
        )

    open_brace = source.find(
        "{",
        match.start(),
        match.end(),
    )

    close_brace = find_matching_delimiter(
        source,
        open_brace,
        "{",
        "}",
    )

    insertion = (
        "\n"
        + "\n".join(
            missing
        )
        + "\n"
    )

    patched = (
        source[:close_brace]
        + insertion
        + source[close_brace:]
    )

    return (
        patched,
        True,
        "added optional rarity/itemType/attunment to Item",
    )


# ---------------------------------------------------------------------------
# FILE DISCOVERY
# ---------------------------------------------------------------------------

def script_folder() -> Path:
    try:
        return Path(
            __file__
        ).resolve().parent
    except NameError:
        return Path.cwd()


def find_default_workbook() -> Optional[Path]:
    for folder in (
        script_folder(),
        Path.cwd(),
    ):
        candidate = (
            folder
            / DEFAULT_WORKBOOK
        )

        if candidate.is_file():
            return candidate

    return None


def find_tribe_files(
    tribes_dir: Path,
    recursive: bool,
) -> List[Path]:
    pattern = (
        "**/tribe*.ts"
        if recursive
        else "tribe*.ts"
    )

    return sorted(
        path
        for path in tribes_dir.glob(
            pattern
        )
        if path.is_file()
        and not path.name.endswith(
            ".before_excel_sync.ts"
        )
    )


def discover_tribes(
    tribe_files: List[Path],
) -> Tuple[
    Dict[str, Path],
    List[str],
]:
    result: Dict[
        str,
        Path,
    ] = {}

    warnings: List[str] = []

    for path in tribe_files:
        try:
            source = path.read_text(
                encoding="utf-8"
            )
        except Exception as error:
            warnings.append(
                f"Could not read {path}: {error}"
            )
            continue

        tribe_name = extract_tribe_name(
            source
        )

        if not tribe_name:
            warnings.append(
                f"Could not determine tribe name in {path.name}"
            )
            continue

        key = normalize_name(
            tribe_name
        )

        if key in result:
            warnings.append(
                f'Duplicate tribe name "{tribe_name}" in '
                f"{result[key].name} and {path.name}"
            )
            continue

        result[key] = path

    return (
        result,
        warnings,
    )


def resolve_tribe_file(
    tribe_display_name: str,
    tribes_dir: Path,
    discovered_by_internal_name: Dict[str, Path],
) -> Tuple[Optional[Path], str]:
    """
    Resolve one Excel tribe to an EXISTING TypeScript file.

    Resolution order:
      1. Explicit TRIBE_FILE_MAP -> exact real filename
      2. Existing file discovered by its internal top-level `name:`
      3. No match -> skip/report

    We intentionally never synthesize a filename from the tribe name.
    """
    mapped_filename = TRIBE_FILE_MAP.get(
        tribe_display_name
    )

    if mapped_filename:
        mapped_path = tribes_dir / mapped_filename

        if mapped_path.is_file():
            return (
                mapped_path,
                "explicit filename map",
            )

        # With --recursive, the exact filename may be in a nested folder.
        recursive_matches = list(
            tribes_dir.glob(
                f"**/{mapped_filename}"
            )
        )

        recursive_matches = [
            path
            for path in recursive_matches
            if path.is_file()
        ]

        if len(recursive_matches) == 1:
            return (
                recursive_matches[0],
                "explicit filename map (nested)",
            )

        if len(recursive_matches) > 1:
            return (
                None,
                (
                    f'ambiguous mapped filename "{mapped_filename}" '
                    f"found in multiple folders"
                ),
            )

        # Don't stop here: the actual file may have been renamed while keeping
        # the correct internal tribe name. Let the safe internal-name fallback
        # try to find it.

    internal_key = normalize_name(
        tribe_display_name
    )

    discovered = discovered_by_internal_name.get(
        internal_key
    )

    if discovered is not None:
        return (
            discovered,
            "existing file internal name",
        )

    if mapped_filename:
        return (
            None,
            (
                f'expected existing file "{mapped_filename}" was not found '
                "and no internal-name fallback matched"
            ),
        )

    return (
        None,
        (
            "no explicit filename mapping and no existing tribe*.ts "
            "with a matching internal name"
        ),
    )


# ---------------------------------------------------------------------------
# SYNC ONE FILE
# ---------------------------------------------------------------------------

def sync_tribe_source(
    source: str,
    excel_items: List[ExcelItem],
    drop_ts_only: bool,
) -> Tuple[
    str,
    List[str],
    Dict[str, int],
]:
    existing_items = parse_existing_items(
        source
    )

    used_existing = set()

    rebuilt: List[str] = []
    report: List[str] = []

    counts = {
        "exact": 0,
        "normalized": 0,
        "prefix": 0,
        "fuzzy": 0,
        "new": 0,
        "ts_only": 0,
        "description_preserved": 0,
        "description_updated": 0,
    }

    for excel_item in excel_items:
        match = find_best_existing_match(
            excel_item,
            existing_items,
            used_existing,
        )

        existing = match.existing

        if existing is not None:
            used_existing.add(
                existing.index
            )

        object_text, desc_reason = (
            build_item_object(
                excel_item,
                existing,
            )
        )

        rebuilt.append(
            object_text
        )

        if match.match_type == "exact":
            counts["exact"] += 1

        elif match.match_type == "parenthetical-normalized":
            counts["normalized"] += 1

        elif match.match_type in {
            "prefix",
            "display-prefix",
        }:
            counts["prefix"] += 1

        elif match.match_type == "fuzzy":
            counts["fuzzy"] += 1

        else:
            counts["new"] += 1

        if "kept existing" in desc_reason:
            counts[
                "description_preserved"
            ] += 1
        else:
            counts[
                "description_updated"
            ] += 1

        if existing is None:
            report.append(
                f'NEW from Excel row {excel_item.row}: '
                f'"{excel_item.name}"'
            )
        else:
            report.append(
                f'{match.match_type.upper()} '
                f'{match.score:.3f}: Excel row {excel_item.row} '
                f'"{excel_item.name}" -> TS "{existing.name}" | '
                f"{desc_reason}"
            )

    # Preserve existing TS-only items unless explicitly asked not to.
    for item in existing_items:
        if item.index in used_existing:
            continue

        counts["ts_only"] += 1

        if not drop_ts_only:
            # Re-indent existing object to fit the new array cleanly without
            # compounding the indentation it already had in the source file.
            raw = textwrap.dedent(
                item.object_text.strip()
            )
            indented = "\n".join(
                "    " + line
                for line in raw.splitlines()
            )
            rebuilt.append(
                indented
            )

        report.append(
            (
                "TS-ONLY PRESERVED: "
                if not drop_ts_only
                else "TS-ONLY DROPPED: "
            )
            + f'"{item.name or "(unnamed item)"}"'
        )

    rewritten = rewrite_items_array(
        source,
        rebuilt,
    )

    return (
        rewritten,
        report,
        counts,
    )


# ---------------------------------------------------------------------------
# MAIN
# ---------------------------------------------------------------------------

def main() -> int:
    parser = argparse.ArgumentParser(
        description=(
            "Sync D&D tribe TypeScript item arrays from "
            '"Put into website as tribes.xlsx".'
        )
    )

    parser.add_argument(
        "workbook",
        nargs="?",
        default=None,
        help=(
            f'Workbook path. Default: "{DEFAULT_WORKBOOK}" beside the script.'
        ),
    )

    parser.add_argument(
        "--sheet",
        default=DEFAULT_SHEET,
        help=(
            f'Excel sheet name. Default: "{DEFAULT_SHEET}"'
        ),
    )

    parser.add_argument(
        "--tribes-dir",
        default=None,
        help=(
            "Folder containing tribe*.ts files. "
            "Default: folder containing this script."
        ),
    )

    parser.add_argument(
        "--recursive",
        action="store_true",
        help=(
            "Search recursively below --tribes-dir for tribe*.ts."
        ),
    )

    parser.add_argument(
        "--dry-run",
        action="store_true",
        help=(
            "Analyze and report without writing TypeScript files."
        ),
    )

    parser.add_argument(
        "--in-place",
        action="store_true",
        help=(
            "Overwrite original tribe files after making "
            "*.before_excel_sync.ts backups."
        ),
    )

    parser.add_argument(
        "--drop-ts-only",
        action="store_true",
        help=(
            "Remove existing TS items that have no spreadsheet match. "
            "Default behavior is to preserve and report them."
        ),
    )

    parser.add_argument(
        "--patch-types",
        action="store_true",
        help=(
            "Patch types.ts Item with optional rarity, itemType, "
            "and attunment fields when types.ts is found."
        ),
    )

    args = parser.parse_args()

    # Resolve workbook.
    if args.workbook:
        workbook_path = Path(
            args.workbook
        ).expanduser().resolve()
    else:
        found = find_default_workbook()

        if found is None:
            print()
            print(
                "ERROR: Could not find:"
            )
            print(
                DEFAULT_WORKBOOK
            )
            print()
            print(
                "Put the workbook beside this Python file "
                "or pass its path explicitly."
            )
            return 1

        workbook_path = found.resolve()

    if not workbook_path.is_file():
        print(
            "ERROR: Workbook not found:"
        )
        print(
            workbook_path
        )
        return 1

    tribes_dir = (
        Path(
            args.tribes_dir
        ).expanduser().resolve()
        if args.tribes_dir
        else script_folder()
    )

    if not tribes_dir.is_dir():
        print(
            "ERROR: Tribe folder not found:"
        )
        print(
            tribes_dir
        )
        return 1

    print()
    print("=" * 80)
    print("D&D TRIBE EXCEL -> TYPESCRIPT SYNC")
    print("=" * 80)
    print(
        "Workbook:",
        workbook_path,
    )
    print(
        "Sheet:   ",
        args.sheet,
    )
    print(
        "Tribes:  ",
        tribes_dir,
    )
    print(
        "Mode:    ",
        (
            "DRY RUN"
            if args.dry_run
            else "IN PLACE"
            if args.in_place
            else "SAFE OUTPUT FOLDER"
        ),
    )
    print()

    try:
        excel_items = read_excel_items(
            workbook_path,
            args.sheet,
        )
    except Exception as error:
        print(
            "ERROR reading workbook:"
        )
        print(
            error
        )
        return 1

    grouped: Dict[
        str,
        List[ExcelItem],
    ] = {}

    tribe_display_names: Dict[
        str,
        str,
    ] = {}

    for item in excel_items:
        key = normalize_name(
            item.tribe_name
        )

        grouped.setdefault(
            key,
            []
        ).append(
            item
        )

        tribe_display_names[
            key
        ] = item.tribe_name

    tribe_files = find_tribe_files(
        tribes_dir,
        args.recursive,
    )

    discovered, discovery_warnings = (
        discover_tribes(
            tribe_files
        )
    )

    print(
        "Spreadsheet items:",
        len(
            excel_items
        ),
    )
    print(
        "Spreadsheet tribes:",
        len(
            grouped
        ),
    )
    print(
        "tribe*.ts files found:",
        len(
            tribe_files
        ),
    )
    print(
        "Tribe names recognized:",
        len(
            discovered
        ),
    )
    print(
        "Explicit real filename mappings:",
        len(TRIBE_FILE_MAP),
    )
    print(
        "New tribe filenames invented: NEVER"
    )
    print()

    output_root = (
        tribes_dir
        if args.in_place
        else (
            tribes_dir
            / OUTPUT_FOLDER_NAME
        )
    )

    if (
        not args.dry_run
        and not args.in_place
    ):
        output_root.mkdir(
            parents=True,
            exist_ok=True,
        )

    overall_counts = {
        "files_updated": 0,
        "exact": 0,
        "normalized": 0,
        "prefix": 0,
        "fuzzy": 0,
        "new": 0,
        "ts_only": 0,
        "description_preserved": 0,
        "description_updated": 0,
    }

    report_lines: List[str] = [
        "D&D Tribe Excel -> TypeScript Sync Report",
        "=" * 80,
        f"Workbook: {workbook_path}",
        f"Sheet: {args.sheet}",
        f"Tribes directory: {tribes_dir}",
        f"Excel items: {len(excel_items)}",
        f"Excel tribe groups: {len(grouped)}",
        f"tribe*.ts files found: {len(tribe_files)}",
        "",
    ]

    if discovery_warnings:
        report_lines.append(
            "FILE DISCOVERY WARNINGS"
        )
        report_lines.append(
            "-" * 80
        )
        report_lines.extend(
            discovery_warnings
        )
        report_lines.append("")

    missing_tribe_files: List[str] = []

    for tribe_key, items in grouped.items():
        tribe_display = tribe_display_names[
            tribe_key
        ]

        ts_path, resolution_method = resolve_tribe_file(
            tribe_display_name=tribe_display,
            tribes_dir=tribes_dir,
            discovered_by_internal_name=discovered,
        )

        if ts_path is None:
            missing_tribe_files.append(
                f"{tribe_display} | {resolution_method}"
            )

            report_lines.extend(
                [
                    f"TRIBE SKIPPED: {tribe_display}",
                    "-" * 80,
                    f"Reason: {resolution_method}",
                    (
                        "No new tribe file was created. "
                        "Add/fix the filename mapping if needed."
                    ),
                    "",
                ]
            )
            continue

        source = ts_path.read_text(
            encoding="utf-8"
        )

        try:
            (
                rewritten,
                tribe_report,
                counts,
            ) = sync_tribe_source(
                source=source,
                excel_items=items,
                drop_ts_only=args.drop_ts_only,
            )
        except Exception as error:
            report_lines.extend(
                [
                    f"TRIBE: {tribe_display}",
                    "-" * 80,
                    f"ERROR processing {ts_path.name}: {error}",
                    "",
                ]
            )
            continue

        overall_counts[
            "files_updated"
        ] += 1

        for key in counts:
            overall_counts[
                key
            ] += counts[
                key
            ]

        report_lines.extend(
            [
                f"TRIBE: {tribe_display}",
                f"File: {ts_path}",
                f"Resolved by: {resolution_method}",
                f"Excel rows: {len(items)}",
                (
                    "Matches: "
                    f"exact={counts['exact']}, "
                    f"normalized={counts['normalized']}, "
                    f"prefix={counts['prefix']}, "
                    f"fuzzy={counts['fuzzy']}, "
                    f"new={counts['new']}, "
                    f"TS-only={counts['ts_only']}"
                ),
                (
                    "Descriptions: "
                    f"format preserved={counts['description_preserved']}, "
                    f"updated from Excel={counts['description_updated']}"
                ),
                "-" * 80,
            ]
        )

        report_lines.extend(
            tribe_report
        )
        report_lines.append("")

        print(
            f"{tribe_display}: "
            f"{len(items)} Excel item(s) | "
            f"{counts['exact'] + counts['normalized'] + counts['prefix'] + counts['fuzzy']} matched | "
            f"{counts['new']} new | "
            f"{counts['ts_only']} TS-only | "
            f"{ts_path.name} [{resolution_method}]"
        )

        if args.dry_run:
            continue

        if args.in_place:
            backup = ts_path.with_name(
                ts_path.stem
                + ".before_excel_sync"
                + ts_path.suffix
            )

            if not backup.exists():
                shutil.copy2(
                    ts_path,
                    backup,
                )

            ts_path.write_text(
                rewritten,
                encoding="utf-8",
            )

        else:
            relative = (
                ts_path.relative_to(
                    tribes_dir
                )
                if args.recursive
                else Path(
                    ts_path.name
                )
            )

            destination = (
                output_root
                / relative
            )

            destination.parent.mkdir(
                parents=True,
                exist_ok=True,
            )

            destination.write_text(
                rewritten,
                encoding="utf-8",
            )

    # Missing tribe TS files.
    if missing_tribe_files:
        report_lines.extend(
            [
                "SPREADSHEET TRIBES SKIPPED - NO EXISTING MATCHING tribe*.ts FILE",
                "-" * 80,
            ]
        )

        for tribe_name in sorted(
            missing_tribe_files
        ):
            report_lines.append(
                tribe_name
            )

        report_lines.append("")

    # Tribe files not represented in Excel.
    excel_keys = set(
        grouped.keys()
    )

    ts_only_tribes = [
        str(path)
        for key, path in discovered.items()
        if key not in excel_keys
    ]

    if ts_only_tribes:
        report_lines.extend(
            [
                "tribe*.ts FILES WITH NO SPREADSHEET TRIBE GROUP",
                "-" * 80,
            ]
        )
        report_lines.extend(
            sorted(
                ts_only_tribes
            )
        )
        report_lines.append("")

    # types.ts handling.
    types_candidates = []

    direct_types = (
        tribes_dir
        / "types.ts"
    )

    if direct_types.is_file():
        types_candidates.append(
            direct_types
        )

    if args.recursive and not types_candidates:
        types_candidates = list(
            tribes_dir.glob(
                "**/types.ts"
            )
        )

    types_note = (
        "types.ts not found"
    )

    if types_candidates:
        types_path = types_candidates[0]

        source = types_path.read_text(
            encoding="utf-8"
        )

        patched, changed, types_note = (
            patch_item_interface(
                source
            )
        )

        if (
            args.patch_types
            and changed
            and not args.dry_run
        ):
            if args.in_place:
                backup = types_path.with_name(
                    types_path.stem
                    + ".before_excel_sync"
                    + types_path.suffix
                )

                if not backup.exists():
                    shutil.copy2(
                        types_path,
                        backup,
                    )

                types_path.write_text(
                    patched,
                    encoding="utf-8",
                )

            else:
                # Copy/patch into safe output folder so imports still work
                # when reviewing the generated tribe files.
                destination = (
                    output_root
                    / types_path.name
                )

                destination.write_text(
                    patched,
                    encoding="utf-8",
                )

        elif (
            not args.in_place
            and not args.dry_run
        ):
            # Copy unchanged types.ts for convenient review.
            destination = (
                output_root
                / types_path.name
            )

            destination.write_text(
                source,
                encoding="utf-8",
            )

    report_lines.extend(
        [
            "TYPES.TS",
            "-" * 80,
            types_note,
            (
                "patch requested: YES"
                if args.patch_types
                else "patch requested: NO"
            ),
            "",
        ]
    )

    report_lines.extend(
        [
            "OVERALL",
            "-" * 80,
            f"Tribe files processed: {overall_counts['files_updated']}",
            f"Exact matches: {overall_counts['exact']}",
            f"Parenthetical-normalized matches: {overall_counts['normalized']}",
            f"Prefix matches: {overall_counts['prefix']}",
            f"Fuzzy matches: {overall_counts['fuzzy']}",
            f"New Excel items added: {overall_counts['new']}",
            f"TS-only items: {overall_counts['ts_only']}",
            f"Descriptions retaining TS formatting: {overall_counts['description_preserved']}",
            f"Descriptions replaced from Excel: {overall_counts['description_updated']}",
            f"Spreadsheet tribes without TS files: {len(missing_tribe_files)}",
            "",
        ]
    )

    report_text = "\n".join(
        report_lines
    )

    print()
    print("=" * 80)
    print("SUMMARY")
    print("=" * 80)
    print(
        "Tribe files processed:",
        overall_counts[
            "files_updated"
        ],
    )
    print(
        "Exact matches:",
        overall_counts[
            "exact"
        ],
    )
    print(
        "Normalized/prefix/fuzzy matches:",
        (
            overall_counts[
                "normalized"
            ]
            + overall_counts[
                "prefix"
            ]
            + overall_counts[
                "fuzzy"
            ]
        ),
    )
    print(
        "New Excel items:",
        overall_counts[
            "new"
        ],
    )
    print(
        "TS-only items:",
        overall_counts[
            "ts_only"
        ],
    )
    print(
        "Spreadsheet tribes missing TS files:",
        len(
            missing_tribe_files
        ),
    )

    if args.dry_run:
        print()
        print(
            "Dry run only. No files were changed."
        )
        print()
        print(
            report_text
        )
        return 0

    report_destination = (
        tribes_dir
        / REPORT_NAME
        if args.in_place
        else output_root
        / REPORT_NAME
    )

    report_destination.write_text(
        report_text,
        encoding="utf-8",
    )

    print()
    if args.in_place:
        print(
            "Original tribe files were updated in place."
        )
        print(
            "Backups use: *.before_excel_sync.ts"
        )
    else:
        print(
            "Updated files:"
        )
        print(
            output_root
        )

    print()
    print(
        "Report:"
    )
    print(
        report_destination
    )

    if not args.patch_types:
        print()
        print(
            "NOTE: If TypeScript reports that rarity/itemType/attunment "
            "are not valid Item fields, rerun with --patch-types "
            "or add those optional fields to interface Item manually."
        )

    return 0


if __name__ == "__main__":
    raise SystemExit(
        main()
    )
